"""
Mind the Context — QA pipeline engine.

Two layers:
  1. Deterministic checks (no API key): length, tags, placeholders, empty.
  2. LLM checks (needs key): register, context-fit, terminology, meaning.

The interesting part is `build_context`: it assembles everything the model needs to
judge a string *in context* — including the exact style-guide section the comment
references and the glossary terms that appear in the source.
"""

from __future__ import annotations

import os
import re
import csv
import json
from dataclasses import dataclass, field, asdict
from typing import Optional

# ---------------------------------------------------------------------------
# Data shapes
# ---------------------------------------------------------------------------

SEVERITY_ORDER = {"error": 0, "warning": 1, "suggestion": 2, "ok": 3}


@dataclass
class Issue:
    severity: str            # error | warning | suggestion
    type: str                # length | tag | placeholder | register | context | terminology | meaning | empty
    message: str
    suggestion: str = ""
    source: str = "rule"     # rule | llm

    def as_dict(self) -> dict:
        return asdict(self)


@dataclass
class StringRow:
    string_id: str
    source: str
    target: str
    language: str
    speaker: Optional[str] = None
    category: Optional[str] = None
    comment: Optional[str] = None
    max_length: Optional[int] = None
    tags: Optional[str] = None


# ---------------------------------------------------------------------------
# Reference asset loading (glossary + style guide)
# ---------------------------------------------------------------------------

_LANG_NORM = {
    "polish": "PL", "pl": "PL", "pol": "PL",
    "german": "DE", "de": "DE", "deu": "DE", "ger": "DE", "deutsch": "DE",
    "english": "EN", "en": "EN", "eng": "EN",
    "russian": "RU", "ru": "RU", "rus": "RU",
    "spanish": "ES", "es": "ES", "spa": "ES",
    "italian": "IT", "it": "IT", "ita": "IT",
    "french": "FR", "fr": "FR", "fra": "FR", "fre": "FR",
    "chinese": "ZH", "zh": "ZH", "korean": "KO", "ko": "KO", "kor": "KO",
    "japanese": "JA", "ja": "JA", "jp": "JA", "jpn": "JA",
    "portuguese": "PT-BR", "pt": "PT", "pt-br": "PT-BR", "ptbr": "PT-BR",
    "dutch": "NL", "nl": "NL", "turkish": "TR", "tr": "TR",
    "ukrainian": "UK", "uk": "UK", "czech": "CS", "cs": "CS",
}


def _norm_lang(s: str) -> str:
    return _LANG_NORM.get((s or "").strip().lower(), (s or "").strip().upper())


def load_glossary(csv_path: str) -> list[dict]:
    """Return glossary rows: [{english, targets: {LANG: term}, category, notes}, ...].
    Supports multiple target-language columns ('Polish Target', 'German Target', or a bare
    language column like 'DE'). The memoQ inflection column is ignored for matching."""
    rows: list[dict] = []
    if not csv_path or not os.path.exists(csv_path):
        return rows
    with open(csv_path, encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        headers = [h for h in (reader.fieldnames or []) if h]

        src_col = next((h for h in headers if "english" in h.lower() or h.strip().lower() in {"source", "src"}), None)
        if src_col is None and headers:
            src_col = headers[0]
        notes_col = next((h for h in headers if h.strip().lower() == "notes"), None)
        cat_col = next((h for h in headers if h.strip().lower() == "category"), None)

        tgt_cols: dict[str, str] = {}
        for h in headers:
            hl = h.strip().lower()
            if h == src_col or hl in {"notes", "category"} or "memoq" in hl or "inflection" in hl:
                continue
            code = None
            if "target" in hl:
                lang_part = re.sub(r"\btarget\b", "", hl).strip(" ()-_")
                code = _norm_lang(lang_part) if lang_part else None
            elif hl in _LANG_NORM:
                code = _LANG_NORM[hl]
            if code:
                tgt_cols.setdefault(code, h)

        for r in reader:
            eng = (r.get(src_col) or "").strip() if src_col else ""
            if not eng:
                continue
            targets = {code: (r.get(h) or "").strip() for code, h in tgt_cols.items() if (r.get(h) or "").strip()}
            rows.append({
                "english": eng,
                "targets": targets,
                "category": (r.get(cat_col) or "").strip() if cat_col else "",
                "notes": (r.get(notes_col) or "").strip() if notes_col else "",
            })
    return rows


def match_glossary(source: str, glossary: list[dict], target_lang: Optional[str] = None, limit: int = 12) -> list[dict]:
    """Glossary entries whose English term appears in the source. When target_lang is given,
    only entries that have a term FOR THAT LANGUAGE are returned — so a Polish-only glossary
    is never enforced against a German (or any other) translation."""
    if not source:
        return []
    low = source.lower()
    code = _norm_lang(target_lang) if target_lang else None
    hits = []
    for g in glossary:
        if not (g["english"] and g["english"].lower() in low):
            continue
        if code is not None:
            term = g["targets"].get(code, "")
            if not term:
                continue  # no term for this language -> don't enforce
        else:
            term = next(iter(g["targets"].values()), "")
        hits.append({"english": g["english"], "term": term, "language": code or "", "notes": g["notes"]})
    hits.sort(key=lambda x: -len(x["english"]))
    return hits[:limit]


_SECTION_RE = re.compile(r"^(#{2,4})\s+(\d+(?:\.\d+)?)\.?\s+(.*)$")


def parse_style_guide(md_path: str) -> dict[str, str]:
    """
    Parse a markdown style guide into {section_number: section_text}.
    Headings look like '## 6. Multi-language utterances' or '### 5.1 Character names'.
    """
    sections: dict[str, str] = {}
    if not md_path or not os.path.exists(md_path):
        return sections
    with open(md_path, encoding="utf-8") as f:
        lines = f.readlines()

    current_num: Optional[str] = None
    buf: list[str] = []

    def flush():
        if current_num is not None:
            sections[current_num] = "".join(buf).strip()

    for line in lines:
        m = _SECTION_RE.match(line.rstrip())
        if m:
            flush()
            current_num = m.group(2)
            buf = [f"§{current_num} {m.group(3)}\n"]
        elif current_num is not None:
            buf.append(line)
    flush()
    return sections


_DOCX_HEADING_RE = re.compile(r"^\s*(\d+(?:\.\d+)?)[.)]?\s+(.+)$")


def parse_style_guide_docx(path: str) -> dict[str, str]:
    """Parse a .docx style guide into {section_number: text}. A paragraph starts a new
    section if it is styled as a heading or looks like a short numbered heading
    (e.g. '6. Multi-language utterances', '5.1 Character names')."""
    sections: dict[str, str] = {}
    if not path or not os.path.exists(path):
        return sections
    try:
        from docx import Document
    except Exception:
        return sections
    doc = Document(path)

    current_num: Optional[str] = None
    buf: list[str] = []

    def flush():
        if current_num is not None:
            sections[current_num] = "\n".join(buf).strip()

    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        style = (getattr(p.style, "name", "") or "").lower()
        is_heading = "heading" in style or "title" in style
        m = _DOCX_HEADING_RE.match(text)
        if m and (is_heading or len(text) < 80):
            flush()
            current_num = m.group(1)
            buf = [f"§{current_num} {m.group(2)}"]
        elif current_num is not None:
            buf.append(text)
    flush()
    return sections


def load_style_guide(path: str) -> dict[str, str]:
    """Dispatch by extension: .docx -> docx parser, otherwise markdown/text parser."""
    if path and path.lower().endswith(".docx"):
        return parse_style_guide_docx(path)
    return parse_style_guide(path)


_REF_RE = re.compile(r"§\s*(\d+(?:\.\d+)?)")


def referenced_sections(comment: Optional[str], sections: dict[str, str], char_budget: int = 1600) -> list[str]:
    """Pull the style-guide sections a comment points to, e.g. 'See §6 and §5.1'."""
    if not comment or not sections:
        return []
    out: list[str] = []
    seen: set[str] = set()
    used = 0
    for num in _REF_RE.findall(comment):
        # try exact, then parent (5.1 -> 5)
        key = num if num in sections else num.split(".")[0]
        if key in sections and key not in seen:
            seen.add(key)
            text = sections[key]
            if len(text) > char_budget:
                text = text[:char_budget] + " …"
            out.append(text)
            used += len(text)
            if used > char_budget * 2:
                break
    return out


# Signal -> keywords to look for in style-guide section headings/bodies. This is matched
# against the ACTUAL guide text (by keyword), not hard-coded section numbers, so it works
# for any style guide that uses descriptive headings.
_CAT_SIGNALS = {
    "ui": ["ui", "button", "length", "label", "menu", "subtitle"],
    "button": ["button", "ui", "length"],
    "dialogue": ["register", "tone", "address", "forms", "dialogue", "profanity", "vocabulary", "period", "names"],
    "item": ["description", "item", "vocabulary", "period", "full stop", "punctuation", "names"],
    "description": ["description", "full stop", "punctuation", "vocabulary"],
    "tooltip": ["ui", "length", "description"],
    "skill": ["skill", "name", "description", "length"],
    "quest": ["quest", "title", "length", "vocabulary"],
    "title": ["title", "length", "names"],
    "location": ["names", "place", "vocabulary", "description"],
    "codex": ["description", "vocabulary", "names"],
    "system": ["placeholder", "gender", "plural", "agreement"],
    "flavor": ["vocabulary", "period", "description"],
    "tutorial": ["placeholder", "ui"],
}
_TAG_SIGNALS = {
    "[fr]": ["french", "multi-language", "language", "tag"],
    "[ru]": ["russian", "multi-language", "language", "tag"],
    "[de]": ["german", "multi-language", "language", "tag"],
    "[lat]": ["latin", "multi-language", "language", "tag"],
    "[ms]": ["gender", "mob", "tag"],
    "[fs]": ["gender", "mob", "tag"],
    "[u]": ["plural", "unit", "number", "count", "placeholder"],
    "{placeholder}": ["placeholder", "tag", "special"],
    "{gender}": ["gender", "agreement", "plural", "placeholder"],
}
# Non-rule sections to skip (overview / Q&A process / quick lists) — heading keywords.
_SKIP_HEADING = ("project overview", "what to flag", "reference quick list")

# Specific signals that should outweigh generic register words, so the directly-relevant
# section (e.g. Multi-language for an [FR] tag) is not crowded out by category signals.
_STRONG = {
    "french", "russian", "german", "latin", "multi-language", "gender", "plural", "unit",
    "count", "number", "mob", "placeholder", "full stop", "punctuation", "length", "button", "title",
}


def auto_sections(row: "StringRow", sections: dict[str, str], budget: int = 2400) -> list[str]:
    """Automatically pick relevant style-guide sections from the string's own signals
    (category, tags, speaker, length, ID) — no comment reference needed."""
    if not sections:
        return []
    signals: set[str] = set()
    cat = (row.category or "").lower()
    for key, kws in _CAT_SIGNALS.items():
        if key in cat:
            signals.update(kws)
    blob = f"{row.tags or ''} {row.source or ''} {row.target or ''}".lower()
    for key, kws in _TAG_SIGNALS.items():
        if key in blob:
            signals.update(kws)
    if row.speaker:
        signals.update(["register", "tone", "address", "forms", "names", "dialogue"])
    if row.max_length:
        signals.update(["ui", "length", "button"])
    sid = (row.string_id or "").lower()
    if "desc" in sid or "description" in cat:
        signals.update(["description", "full stop", "punctuation"])
    if not signals:
        return []

    scored = []
    for num, text in sections.items():
        head = text.split("\n", 1)[0].lower()
        if any(s in head for s in _SKIP_HEADING):
            continue
        body = text.lower()
        score = 0
        for kw in signals:
            if kw not in body:
                continue
            w = 3 if kw in _STRONG else 1
            score += (2 * w) if kw in head else w
        if score > 0:
            scored.append((score, num, text))
    scored.sort(key=lambda x: (-x[0], x[1]))

    out, used = [], 0
    for _, _, text in scored:
        t = text if len(text) <= 1000 else text[:1000] + " …"
        if used + len(t) > budget and out:
            break
        out.append(t)
        used += len(t)
    return out


# ---------------------------------------------------------------------------
# Character bible (matched by Speaker — for dialogue strings)
# ---------------------------------------------------------------------------

# Titles/particles to ignore when matching a Speaker to a bible entry.
_TITLE_STOP = {
    "dr", "mr", "mrs", "ms", "count", "countess", "captain", "capt", "lt",
    "lieutenant", "father", "grand", "duke", "sister", "general", "senator",
    "pan", "pani", "panna", "herbu", "the", "of", "st", "von",
}


def load_character_bible(md_path: str) -> list[dict]:
    """Parse a character bible markdown into [{name, voice}, ...].
    Each '### Name' heading starts one character block."""
    chars: list[dict] = []
    if not md_path or not os.path.exists(md_path):
        return chars
    with open(md_path, encoding="utf-8") as f:
        lines = f.readlines()

    name: Optional[str] = None
    buf: list[str] = []

    def flush():
        if name:
            block = "".join(buf).strip()
            chars.append({"name": name, "voice": _voice_snippet(block)})

    for line in lines:
        m = re.match(r"^###\s+(.*?)\s*$", line.rstrip())
        if m:
            flush()
            name = m.group(1).strip()
            buf = []
        elif name is not None:
            buf.append(line)
    flush()
    return chars


def _voice_snippet(block: str, budget: int = 900) -> str:
    """Prefer the 'Voice and personality' paragraph; fall back to the block start."""
    low = block.lower()
    key = "voice and personality"
    i = low.find(key)
    if i != -1:
        rest = block[i + len(key):]
        cut = re.search(r"\n\s*\n\*\*", rest)          # stop at next bold sub-heading
        snippet = rest[: cut.start()] if cut else rest
        snippet = snippet.lstrip(".:* \n")
        return snippet.strip()[:budget]
    return block[:budget]


def _name_tokens(s: str) -> set:
    toks = re.findall(r"[^\W\d_]+", (s or "").lower(), re.UNICODE)
    return {t for t in toks if len(t) > 2 and t not in _TITLE_STOP}


def match_character(speaker: Optional[str], characters: list[dict]) -> Optional[dict]:
    """Best character-bible entry for a Speaker, by name-token overlap."""
    if not speaker or not characters:
        return None
    want = _name_tokens(speaker)
    if not want:
        return None
    best, best_score = None, 0
    for c in characters:
        score = len(want & _name_tokens(c["name"]))
        if score > best_score:
            best, best_score = c, score
    return best if best_score >= 1 else None


def match_entities(source: str, characters: list[dict], exclude: Optional[dict] = None,
                   limit: int = 6) -> list[dict]:
    """Character/entity profiles whose name appears in the SOURCE text.
    Covers descriptive (non-dialogue) segments where a named entity (e.g. a skill or
    character) is mentioned but is not the speaker — so its facts (gender, etc.) still apply."""
    if not source or not characters:
        return []
    src_tokens = _name_tokens(source)
    src_low = source.lower()
    excl_name = exclude["name"] if exclude else None
    hits: list[dict] = []
    for c in characters:
        if c["name"] == excl_name:
            continue
        ntoks = _name_tokens(c["name"])
        if ntoks and (ntoks & src_tokens or c["name"].lower() in src_low):
            hits.append(c)
    return hits[:limit]


def parse_character_text(text: str) -> list[dict]:
    """Parse free-pasted character descriptions into [{name, voice}, ...].
    Tolerant of several formats, so the user can paste from any source (Word, PDF, wiki):
      - '### Name' or '**Name**' headings
      - 'Name: description' / 'Name — description' lines
      - blank-line-separated paragraphs that begin with the character's name
    """
    if not text or not text.strip():
        return []
    text = text.replace("\r\n", "\n").strip()

    # 1) Heading-delimited (markdown # or bold **Name**)
    if re.search(r"(?m)^\s*(#{1,6}\s+\S|\*\*[^*\n]+\*\*\s*:?\s*$)", text):
        return _parse_headed_blocks(text)

    # 2) Otherwise split into blocks; a block whose every line is "Name<delim> desc"
    #    (or "**Name** desc") is treated as a list — one entry per line.
    out: list[dict] = []
    for block in re.split(r"\n\s*\n", text):
        lines = [ln for ln in block.split("\n") if ln.strip()]
        if not lines:
            continue
        line_entries = [_line_name_desc(ln) for ln in lines]
        if len(lines) > 1 and all(line_entries):
            for nm, ds in line_entries:
                out.append({"name": nm, "voice": _voice_snippet(ds)})
        elif len(lines) == 1 and line_entries[0]:
            nm, ds = line_entries[0]
            out.append({"name": nm, "voice": _voice_snippet(ds)})
        else:
            name, desc = _split_name_desc(block)
            if name:
                out.append({"name": name, "voice": _voice_snippet(desc)})
    return out


def _line_name_desc(line: str):
    """Parse a single line of the form '**Name** desc', 'Name: desc', or 'Name — desc'."""
    line = line.strip()
    m = re.match(r"^\*\*\s*([^*\n]+?)\s*\*\*\s*[:—–-]?\s*(.+)$", line)   # **Name** desc
    if m and _looks_like_name(m.group(1)):
        return m.group(1).strip(), m.group(2).strip()
    m = re.match(r"^(.{1,40}?)\s*[—–:]\s+(.+)$", line) or re.match(r"^(.{1,40}?)\s+-\s+(.+)$", line)
    if m and _looks_like_name(m.group(1)):
        return m.group(1).strip(), m.group(2).strip()
    return None


def _parse_headed_blocks(text: str) -> list[dict]:
    out: list[dict] = []
    name: Optional[str] = None
    buf: list[str] = []

    def flush():
        if name:
            out.append({"name": name, "voice": _voice_snippet("".join(buf).strip())})

    for line in text.split("\n"):
        m = re.match(r"^\s*(?:#{1,6}\s+(.*\S)|\*\*([^*\n]+?)\*\*)\s*:?\s*$", line)
        if m:
            flush()
            name = (m.group(1) or m.group(2)).strip()
            buf = []
        elif name is not None:
            buf.append(line + "\n")
    flush()
    return out


def _looks_like_name(s: str) -> bool:
    s = s.strip().lstrip("*").strip()
    return bool(s) and s[0].isupper() and 1 <= len(s.split()) <= 4


def _leading_name(line: str) -> Optional[str]:
    """Leading run of capitalized words at the start of a line (unicode-aware)."""
    words = line.split()
    name_words: list[str] = []
    for w in words[:5]:
        core = w.strip(".,;:—–-()")
        if core and core[0].isupper():
            name_words.append(core)
        else:
            break
    return " ".join(name_words) if name_words else None


def _split_name_desc(block: str):
    """From one pasted block, separate the character name from the description."""
    first, _, rest = block.partition("\n")
    # delimiter forms: "Name: ...", "Name — ...", "Name - ..."
    m = re.match(r"^\s*(.{1,40}?)\s*[—–:]\s+(.*)$", first) or re.match(r"^\s*(.{1,40}?)\s+-\s+(.*)$", first)
    if m and _looks_like_name(m.group(1)):
        desc = (m.group(2) + ("\n" + rest if rest else "")).strip()
        return m.group(1).strip(), desc
    # otherwise: leading capitalized words are the name, the whole block is the description
    name = _leading_name(first)
    if name:
        return name, block.strip()
    return None, block


# ---------------------------------------------------------------------------
# Deterministic checks (no API key required)
# ---------------------------------------------------------------------------

_TAG_RE = re.compile(r"\[/?[A-Za-z]{1,8}\]")           # [FR] [/FR] [LAT] [u] [one] [few] [ms] [fs] …
_PLACEHOLDER_RE = re.compile(r"\{[a-zA-Z0-9_.\-]+\}")  # {placeholder} {gender} {QUANTITY} {character.firstName}


def deterministic_checks(row: StringRow) -> list[Issue]:
    issues: list[Issue] = []
    src = row.source or ""
    tgt = (row.target or "").strip()

    if not tgt:
        issues.append(Issue("warning", "empty", f"No {row.language} translation provided."))
        return issues  # nothing else to check

    # 1. Max length
    if row.max_length and int(row.max_length) > 0 and len(tgt) > int(row.max_length):
        over = len(tgt) - int(row.max_length)
        issues.append(
            Issue(
                "error",
                "length",
                f"Translation is {len(tgt)} chars; limit is {row.max_length} (over by {over}).",
                suggestion="Shorten to fit the UI element this string belongs to.",
            )
        )

    # 2. Tag preservation ([FR]...[/FR], [RU], [LAT], …)
    src_tags = _TAG_RE.findall(src)
    tgt_tags = _TAG_RE.findall(tgt)
    for tag in set(src_tags):
        if src_tags.count(tag) > tgt_tags.count(tag):
            issues.append(
                Issue(
                    "error",
                    "tag",
                    f"Tag {tag} present in source but missing/under-count in translation.",
                    suggestion=f"Keep {tag} exactly as in the source.",
                )
            )

    # 3. Placeholder preservation ({placeholder}, {gender})
    src_ph = _PLACEHOLDER_RE.findall(src)
    tgt_ph = _PLACEHOLDER_RE.findall(tgt)
    for ph in set(src_ph):
        if src_ph.count(ph) > tgt_ph.count(ph):
            issues.append(
                Issue(
                    "error",
                    "placeholder",
                    f"Placeholder {ph} is missing from the translation.",
                    suggestion=f"Reinsert {ph}; the engine substitutes it at runtime.",
                )
            )

    # 4. Mob gender-variant nouns must carry both forms (style guide §6.7)
    has_ms, has_fs = "[ms]" in tgt, "[fs]" in tgt
    if has_ms != has_fs:
        missing = "[fs]" if has_ms else "[ms]"
        issues.append(
            Issue(
                "error",
                "tag",
                f"Mob noun has only one gender form — missing the {missing} variant (style guide §6.7 requires both).",
                suggestion="Provide both forms, e.g. [ms]arystokrata;[fs]arystokratka.",
            )
        )

    # 5. Descriptions must end with a full stop (style guide §11)
    if _is_description(row) and not _ends_with_terminal_punct(tgt):
        issues.append(
            Issue(
                "warning",
                "punctuation",
                "Description does not end with a full stop (style guide: descriptions end with full stops).",
                suggestion="End the description with a period.",
            )
        )

    return issues


def _is_description(row: StringRow) -> bool:
    sid = (row.string_id or "").lower()
    cat = (row.category or "").lower()
    return "description" in cat or bool(re.search(r"(^|[_\-.])desc(\b|[_\-.]|$)", sid))


# trailing closing quotes / brackets / parens, and trailing tags/placeholders, to look past
_TRAIL_CLOSERS = r'[\s"\'`»«”“„’‚)\]\}>]+'


def _ends_with_terminal_punct(tgt: str) -> bool:
    """True if the visible text ends with sentence-ending punctuation, ignoring trailing
    quotes/brackets and any trailing tags/placeholders like [/u] or {QUANTITY}."""
    s = (tgt or "").rstrip()
    if not s:
        return True  # empty handled elsewhere; don't double-flag
    prev = None
    while s and s != prev:
        prev = s
        s = re.sub(_TRAIL_CLOSERS + r"$", "", s)
        s = re.sub(r"(?:\[/?[A-Za-z]{1,8}\]|\{[^}]*\})+$", "", s).rstrip()
    return bool(s) and s[-1] in ".!?…"


# ---------------------------------------------------------------------------
# Context assembly (shared by the LLM layer; also nice to display in the UI)
# ---------------------------------------------------------------------------

def build_context(row: StringRow, glossary: list[dict], style_sections: dict[str, str],
                  characters: Optional[list[dict]] = None) -> dict:
    gloss_hits = match_glossary(row.source, glossary, row.language)
    # Comment references take priority, then auto-selected sections by the string's signals.
    ref = referenced_sections(row.comment, style_sections)
    auto = auto_sections(row, style_sections)
    style_hits, seen, used = [], set(), 0
    for s in ref + auto:
        key = s.split("\n", 1)[0]
        if key in seen:
            continue
        if used + len(s) > 3400 and style_hits:
            break
        seen.add(key)
        style_hits.append(s)
        used += len(s)
    inferred = _infer_from_id(row)
    character = match_character(row.speaker, characters or [])
    entities = match_entities(row.source, characters or [], exclude=character)
    return {
        "inferred": inferred,
        "glossary": gloss_hits,
        "style_sections": style_hits,
        "character": character,
        "entities": entities,
    }


def _infer_from_id(row: StringRow) -> list[str]:
    """Cheap heuristics that mirror what the model is asked to reason about."""
    notes: list[str] = []
    sid = (row.string_id or "").lower()
    cat = (row.category or "").lower()
    if sid.startswith("ui_") or cat == "ui":
        notes.append("UI element — likely space-constrained; prefer concise, conventional wording.")
    if "button" in (row.comment or "").lower():
        notes.append("Button label — imperative, no trailing punctuation.")
    if sid.startswith("dlg_") or cat == "dialogue":
        notes.append("Spoken dialogue — register must match the speaker and scene.")
    if "tooltip" in cat:
        notes.append("Tooltip — terse, informative.")
    if row.max_length:
        notes.append(f"Hard limit {row.max_length} characters.")
    if row.speaker:
        notes.append(f"Speaker: {row.speaker} — match their voice (see character bible).")
    return notes


# ---------------------------------------------------------------------------
# LLM layer
# ---------------------------------------------------------------------------

LLM_SYSTEM = """You are a senior localization QA reviewer. You judge whether a translation is \
correct *in its context*, not just in isolation. You are given a source string plus rich context: \
the string ID, category, max length, tags, the speaker, a translator comment, the relevant \
style-guide sections, and matching glossary terms.

Report only real problems. For each, return an object with:
  severity: "error" | "warning" | "suggestion"
  type: "register" | "context" | "terminology" | "meaning" | "tag" | "length" | "other"
  message: one concrete sentence describing the problem
  suggestion: a brief concrete fix (may be empty)

Rules:
- A comment is binding. If it says "question form, NOT a button label", a button-style \
translation is an error.
- Glossary terms must be used in the prescribed form.
- Judge every translation by the grammar of ITS OWN target language. NEVER apply a rule from one \
language to another where the grammatical feature does not exist. Style-guide sections may be written \
for one language (e.g. Polish); apply such language-specific rules only to that language.
- Grammatical GENDER rules apply ONLY in target languages whose verbs/adjectives inflect for gender \
(e.g. Polish, Russian, Ukrainian). In German, English, etc. past-tense verbs and participles are NOT \
gendered — do NOT raise any gender issue there.
  - Fixed-gender referent (in a gendered language): if a named entity is male but the translation uses \
feminine forms (or vice versa), that is an error.
  - Variable-gender referent (in a gendered language): a player-name placeholder like \
{character.firstName} can be male or female, so the form must be gender-NEUTRAL — flag a gendered verb/\
adjective or a bracketed "(a)" form. In a non-gendered language (German/English) this is a non-issue.
- If the translation is good, return an empty list.
Return ONLY valid JSON: {"issues": [ ... ]}. No prose."""


def _build_llm_user_prompt(row: StringRow, ctx: dict) -> str:
    lines = [
        f"Target language: {row.language}",
        f"String ID: {row.string_id}",
        f"Category: {row.category or '—'}",
        f"Speaker: {row.speaker or '—'}",
        f"Max length: {row.max_length or 'unconstrained'}",
        f"Tags: {row.tags or '—'}",
        f"Translator comment: {row.comment or '(none — infer context from the ID and metadata)'}",
        "",
        f"SOURCE (EN): {row.source}",
        f"TRANSLATION ({row.language}): {row.target}",
    ]
    if ctx["inferred"]:
        lines += ["", "Inferred context:"] + [f"- {n}" for n in ctx["inferred"]]
    if ctx["glossary"]:
        lines += ["", f"Relevant glossary terms — the required {row.language} form (English -> {row.language}). "
                  "Apply ONLY to this target language:"]
        lines += [f"- {g['english']} -> {g['term']}" + (f" ({g['notes']})" if g["notes"] else "") for g in ctx["glossary"]]
    if ctx["style_sections"]:
        lines += ["", "Relevant style-guide sections:"]
        lines += [s for s in ctx["style_sections"]]
    if ctx.get("character"):
        ch = ctx["character"]
        lines += ["", f"Speaker profile — {ch['name']} (the translation must match this voice):", ch["voice"]]
    if ctx.get("entities"):
        lines += ["", "Named entities mentioned in the source \u2014 respect these facts "
                  "(including grammatical gender agreement in the target language):"]
        lines += [f"- {e['name']}: {e['voice'][:280]}" for e in ctx["entities"]]
    return "\n".join(lines)


def _chat(client, model: str, messages: list[dict]):
    """One chat completion. Tries JSON mode first (cloud models), falls back without it
    for local servers (Ollama / LM Studio) that may not support response_format.
    Returns (text, usage) where usage = {"prompt": int, "completion": int}."""
    try:
        resp = _create(client, model, messages)
    except UnicodeEncodeError:
        # The runtime cannot encode non-ASCII (e.g. a C/ascii locale): retry with the message
        # content escaped to ASCII so the request still goes through.
        safe = [
            {"role": m["role"], "content": str(m["content"]).encode("ascii", "backslashreplace").decode("ascii")}
            for m in messages
        ]
        resp = _create(client, model, safe)
    text = resp.choices[0].message.content or ""
    u = getattr(resp, "usage", None)
    usage = {
        "prompt": int(getattr(u, "prompt_tokens", 0) or 0),
        "completion": int(getattr(u, "completion_tokens", 0) or 0),
    }
    return text, usage


def _create(client, model: str, messages: list[dict]):
    """One create() call: JSON mode first, fall back without it for servers that reject it.
    A UnicodeEncodeError is re-raised (handled by the caller), not swallowed by the fallback."""
    try:
        return client.chat.completions.create(
            model=model, messages=messages, max_tokens=900, temperature=0,
            response_format={"type": "json_object"},
        )
    except UnicodeEncodeError:
        raise
    except Exception:
        return client.chat.completions.create(
            model=model, messages=messages, max_tokens=900, temperature=0,
        )


def llm_checks(row: StringRow, ctx: dict, client, model: str):
    """Call the model. `client` is an OpenAI-compatible instance (cloud or local).
    Returns (issues, usage) — usage = {"prompt": int, "completion": int, "calls": int}."""
    zero = {"prompt": 0, "completion": 0, "calls": 0}
    if client is None or not (row.target or "").strip():
        return [], zero
    messages = [
        {"role": "system", "content": LLM_SYSTEM},
        {"role": "user", "content": _build_llm_user_prompt(row, ctx)},
    ]
    try:
        text, u = _chat(client, model, messages)
        usage = {"prompt": u["prompt"], "completion": u["completion"], "calls": 1}
        data = _extract_json(text)
        out: list[Issue] = []
        for it in data.get("issues", []):
            out.append(
                Issue(
                    severity=str(it.get("severity", "warning")).lower(),
                    type=str(it.get("type", "other")).lower(),
                    message=str(it.get("message", "")).strip(),
                    suggestion=str(it.get("suggestion", "")).strip(),
                    source="llm",
                )
            )
        return [i for i in out if i.message], usage
    except Exception as e:  # never let one bad row kill the run
        return [Issue("warning", "other", f"LLM check skipped: {e}", source="llm")], zero


def _extract_json(text: str) -> dict:
    text = text.strip()
    if text.startswith("```"):
        text = re.sub(r"^```[a-zA-Z]*\n?|\n?```$", "", text).strip()
    try:
        return json.loads(text)
    except Exception:
        m = re.search(r"\{.*\}", text, re.DOTALL)
        if m:
            try:
                return json.loads(m.group(0))
            except Exception:
                return {}
        return {}


# ---------------------------------------------------------------------------
# Orchestration
# ---------------------------------------------------------------------------

def analyze_row(row: StringRow, glossary, style_sections, client=None, model: str = "gpt-4o",
                use_llm: bool = True, characters=None) -> dict:
    ctx = build_context(row, glossary, style_sections, characters)
    issues = deterministic_checks(row)
    tokens = {"prompt": 0, "completion": 0, "calls": 0}
    if use_llm and client is not None:
        llm_issues, tokens = llm_checks(row, ctx, client, model)
        issues += llm_issues

    issues.sort(key=lambda i: SEVERITY_ORDER.get(i.severity, 9))
    worst = issues[0].severity if issues else "ok"
    return {
        "string_id": row.string_id,
        "language": row.language,
        "source": row.source,
        "target": row.target,
        "status": worst,
        "n_issues": len(issues),
        "issues": [i.as_dict() for i in issues],
        "context": ctx,
        "tokens": tokens,
    }
